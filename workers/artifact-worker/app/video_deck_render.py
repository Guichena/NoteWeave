from __future__ import annotations

from io import BytesIO
import math
from pathlib import Path
import re
import unicodedata

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from app.config import resolve_mcp_sandbox_root


WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)
IMAGE_BOX = (Inches(0.7), Inches(1.45), Inches(7.65), Inches(5.25))
TEXT_BOX = (Inches(8.7), Inches(1.45), Inches(3.9), Inches(4.9))


def render_original_video_deck(task_id: str, ir: dict, read_file=None, *, evidence=None) -> Path:
    """读取 DECK 并嵌入原帧，保持文本可编辑。"""
    return _render_deck(task_id, ir, evidence, read_file)


def _display_units(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
               for char in text)


def _title_font_size(title: str) -> int:
    units = _display_units(title)
    if units > 80:
        raise ValueError("video deck title exceeds fixed template width")
    return 30 if units <= 45 else 24 if units <= 70 else 20


def _body_font_size(claims: list[str]) -> int:
    text = claims or [""]
    width_pt = TEXT_BOX[2] / 12700
    height_pt = TEXT_BOX[3] / 12700
    for size in (18, 16, 14):
        lines = sum(max(1, math.ceil(_display_units(claim) * size * 0.6 / width_pt))
                    for claim in text) + 2 * (len(text) - 1)
        if lines * size * 1.35 <= height_pt:
            return size
    raise ValueError("video deck claims exceed fixed template height")


def _fit_image(width: int, height: int) -> tuple[int, int, int, int]:
    left, top, box_width, box_height = IMAGE_BOX
    scale = min(box_width / width, box_height / height)
    fitted_width = round(width * scale)
    fitted_height = round(height * scale)
    return (left + (box_width - fitted_width) // 2,
            top + (box_height - fitted_height) // 2,
            fitted_width, fitted_height)


def _render_deck(task_id, ir, evidence, frame_reader):
    from app.pipeline.renderers.frames import read_frame, time_label
    if ir.get("kind") != "DECK" or not re.fullmatch(r"[A-Za-z0-9_-]{1,128}", task_id):
        raise ValueError("幻灯片形态或任务编号无效")
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = WIDTH, HEIGHT
    images = []
    for index, spec in enumerate(ir["slides"], 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(250, 249, 246)
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.42), Inches(11.9), Inches(0.7))
        title.text_frame.text = spec["title"]
        p = title.text_frame.paragraphs[0]
        p.font.name, p.font.size, p.font.bold = "Noto Sans CJK SC", Pt(_title_font_size(spec["title"])), True
        p.font.color.rgb = RGBColor(25, 44, 62)
        content = None
        figure = spec.get("figure")
        if figure:
            content = read_frame(figure["frame_id"], evidence, frame_reader)
            with Image.open(BytesIO(content)) as image:
                box = _fit_image(*image.size)
            slide.shapes.add_picture(BytesIO(content), *box[:2], width=box[2], height=box[3])
        images.append(content)
        claims = [bullet["text"] for bullet in spec["bullets"]]
        body = slide.shapes.add_textbox(*(TEXT_BOX if figure else (Inches(1.1), Inches(1.6), Inches(11.1), Inches(4.7))))
        body.text_frame.word_wrap = True
        body.text_frame.text = "\n\n".join(claims)
        # 无配图的页面仍按较窄栏宽估算，避免文本溢出。
        size = _body_font_size(claims)
        for paragraph in body.text_frame.paragraphs:
            paragraph.font.name, paragraph.font.size = "Noto Sans CJK SC", Pt(size)
            paragraph.font.color.rgb = RGBColor(38, 51, 60)
        footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.9), Inches(11.9), Inches(0.35))
        footer.text_frame.text = (time_label(figure, evidence) + " · " if figure else "") + f"{index}/{len(ir['slides'])}"
        footer.text_frame.paragraphs[0].font.size = Pt(10)
        slide.notes_slide.notes_text_frame.text = spec.get("speaker_notes", {}).get("text", "")
    output = resolve_mcp_sandbox_root() / "bilibili-render-pdf" / "exports" / task_id
    output.mkdir(parents=True, exist_ok=True)
    path = output / "learning-deck.pptx"
    presentation.save(path)
    saved = Presentation(path)
    if len(saved.slides) != len(ir["slides"]):
        raise ValueError("幻灯片保存后页数与 IR 不一致")
    for slide, spec, content in zip(saved.slides, ir["slides"], images):
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if len(pictures) != (1 if content else 0) or (content and (pictures[0].image.blob != content
                or any((pictures[0].crop_left, pictures[0].crop_right, pictures[0].crop_top, pictures[0].crop_bottom)))):
            raise ValueError("幻灯片改变或裁剪了原始关键帧")
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        if spec["title"] not in text or any(b["text"] not in text for b in spec["bullets"]):
            raise ValueError("幻灯片丢失可编辑正文")
    return path
