"""覆盖 DECK 规划、原帧校验和实际文档渲染。"""

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import shutil

from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
import pytest

from app.pipeline.evidence import evidence_errors
from app.pipeline.renderers import pdf
from app.pipeline.renderers.frames import read_frame
from app.pipeline.rules import contract_errors, validate_contract
from app.pipeline.runner import run
from app.pipeline.schemas import load_schema, schema_errors
from app.pipeline.skills import SkillCatalog
from app.pipeline.stages import check, plan
from app.video_deck_render import render_original_video_deck
from test_pipeline import Model, task


FACT = {"text": "缓存有效期为30秒。", "anchors": ["s0001"]}
EVIDENCE = {"sentences": [{"id": "s0001", "text": FACT["text"], "source_id": "source",
    "source_snapshot_id": "snapshot", "chunk_id": "chunk", "frame_ids": ["frame"],
    "part": 2, "start_ms": 1000, "end_ms": 9000}]}
FIGURE = {"frame_id": "frame", "caption": FACT}
DECK = {"schema_version": 1, "kind": "DECK", "title": "缓存原理",
        "slides": [{"id": "u0", "title": "缓存有效期", "bullets": [FACT], "figure": FIGURE,
                    "speaker_notes": FACT}]}
DOC = {"schema_version": 1, "kind": "DOC", "title": "缓存讲义", "sections": [{"id": "u0", "heading": "缓存有效期",
    "blocks": [{"id": "p0", "type": "paragraph", "statements": [{"id": "s0", **FACT}]},
               {"id": "f0", "type": "figure", **FIGURE}]}]}


def image_bytes():
    image = Image.new("RGB", (960, 540), "#caddeb")
    from PIL import ImageDraw
    draw = ImageDraw.Draw(image)
    draw.rectangle((70, 80, 890, 460), fill="white", outline="#244e70", width=4)
    draw.text((140, 210), "Cache -> 30s -> Reload", fill="#244e70", font_size=42)
    stream = BytesIO()
    image.save(stream, format="PNG")
    return stream.getvalue()


class DeckModel(Model):
    def __init__(self, kind, with_frames):
        super().__init__(kind)
        self.with_frames = with_frames

    def complete_json(self, purpose, payload):
        stage, data = purpose.rsplit(".", 1)[-1], payload["inputs"]
        if stage == "plan":
            response = json.loads(super().complete_json(purpose, payload))
            if self.with_frames:
                for unit in response["units"]:
                    unit["frame_ids"] = ["frame"]
            return json.dumps(response)
        if stage == "draft":
            unit = data["unit"]
            item = {"id": unit["id"], "title": unit["heading"], "bullets": [FACT], "speaker_notes": FACT} if self.kind == "DECK" else {
                "id": unit["id"], "heading": unit["heading"], "blocks": [{"id": "p0", "type": "paragraph", "statements": [{"id": "s0", **FACT}]}]}
            if self.with_frames:
                if self.kind == "DECK":
                    item["figure"] = FIGURE
                else:
                    item["blocks"].append({"id": "f0", "type": "figure", **FIGURE})
            self.items[unit["id"]] = item
            return json.dumps({"item": item})
        return super().complete_json(purpose, payload)


def test_deck_schema_and_figure_contract_reject_missing_or_unrelated_frames():
    assert not schema_errors(DECK, load_schema("DECK"))
    bad = deepcopy(DECK)
    del bad["slides"][0]["figure"]
    assert contract_errors(bad, {"figure_required": True})
    bad = deepcopy(DECK)
    bad["slides"][0]["figure"]["frame_id"] = "foreign"
    assert evidence_errors(bad, EVIDENCE)
    assert contract_errors(DECK, {"deck_bullets": {"min": 2, "max": 6}})
    for kind in ["DOC", "DECK"]:
        validate_contract(kind, {"figure_required": ["u0"]})
    with pytest.raises(ValueError):
        validate_contract("TABLE", {"figure_required": True})


def test_plan_requires_frames_from_the_units_own_anchors():
    catalog = SkillCatalog()
    source = task("video-deck", catalog)
    with pytest.raises(ValueError, match="关键帧"):
        plan.plan(source, catalog.resolve("video-deck"), {"count": "5"}, EVIDENCE, DeckModel("DECK", False))
    with pytest.raises(ValueError, match="不属于"):
        plan.plan(source, catalog.resolve("video-deck"), {"count": "5"}, {"sentences": [{**EVIDENCE["sentences"][0], "frame_ids": []}]}, DeckModel("DECK", True))


def test_frame_decode_scope_and_caption_are_checked():
    assert read_frame("frame", EVIDENCE, lambda _: image_bytes()) == image_bytes()
    with pytest.raises(ValueError, match="素材包"):
        read_frame("foreign", EVIDENCE, lambda _: image_bytes())
    with pytest.raises(ValueError, match="读取接口"):
        read_frame("frame", EVIDENCE, None)
    bad_evidence = {"sentences": [*EVIDENCE["sentences"], {"id": "s0002", "text": FACT["text"], "frame_ids": []}]}
    bad = deepcopy(DECK)
    bad["slides"][0]["figure"]["caption"] = {**FACT, "anchors": ["s0002"]}
    assert "图注引用的素材未包含该关键帧" in evidence_errors(bad, bad_evidence)


def test_callback_reads_binary_frame_with_current_delivery_token(monkeypatch):
    from email.message import Message
    from app.callback import JavaArtifactCallbackClient
    content = image_bytes()
    seen = []
    class Response(BytesIO):
        headers = Message()
    Response.headers["Content-Type"] = "image/png"
    def open_request(req, timeout):
        seen.append(req)
        return Response(content)
    monkeypatch.setattr("app.callback.credential_safe_urlopen", open_request)
    client = JavaArtifactCallbackClient("http://host", internal_auth_token="internal", delivery_token="delivery")
    assert client.read_artifact_frame("task/id", "frame/id") == content
    assert seen[0].full_url == "http://host/internal/worker/artifact-tasks/task%2Fid/frames/frame%2Fid"
    assert dict((k.lower(), v) for k, v in seen[0].header_items())["x-noteweave-outbox-delivery-token"] == "delivery"
    Response.headers.replace_header("Content-Type", "text/html")
    with pytest.raises(ValueError, match="图片类型"):
        client.read_artifact_frame("task", "frame")


def test_pptx_keeps_original_frame_and_editable_text_and_supports_plain_pages(monkeypatch, tmp_path):
    monkeypatch.setenv("NOTEWEAVE_MCP_SANDBOX_ROOT", str(tmp_path))
    ir = deepcopy(DECK)
    ir["slides"].append({"id": "u1", "title": "普通文字页", "bullets": [FACT]})
    original = image_bytes()
    path = render_original_video_deck("deck-test", ir, evidence=EVIDENCE, read_file=lambda _: original)
    saved = Presentation(path)
    assert len(saved.slides) == 2
    picture = next(s for s in saved.slides[0].shapes if s.shape_type == 13)
    assert picture.image.blob == original
    assert not any([picture.crop_left, picture.crop_right, picture.crop_top, picture.crop_bottom])
    assert not any(s.shape_type == 13 for s in saved.slides[1].shapes)
    assert FACT["text"] in saved.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.skipif(shutil.which("xelatex") is None, reason="运行环境未安装 XeLaTeX")
def test_real_pdf_uses_template_original_frame_and_time_interval(tmp_path):
    original = image_bytes()
    path = pdf.render(DOC, tmp_path, EVIDENCE, lambda _: original)
    text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    # Fandol 的 CID 字体依赖阅读器的 Adobe-GB1 映射，正文以渲染图另行核对。
    assert r"\newcommand{\notetitle}{缓存讲义}" in (tmp_path / "lecture-notes.tex").read_text(encoding="utf-8")
    assert "00:00:01" in text and "00:00:09" in text
    assert next(tmp_path.glob("frame-*.png")).read_bytes() == original
    assert "在此填写" not in text and "请在生成" not in text
    assert pdf.escape(r"x_%\\input{evil}") == r"x\_\%\textbackslash{}\textbackslash{}input\{evil\}"


@pytest.mark.parametrize("name", ["slide-deck", "video-deck", "lecture-notes-pdf"])
def test_new_skills_share_all_five_stages_and_produce_downloadable_files(name, monkeypatch, tmp_path):
    if shutil.which("xelatex" if name == "lecture-notes-pdf" else "soffice") is None:
        pytest.skip("运行环境未安装对应文档渲染工具")
    monkeypatch.setenv("NOTEWEAVE_MCP_SANDBOX_ROOT", str(tmp_path))
    catalog = SkillCatalog()
    source = task(name, catalog)
    def search(task_id, query, limit):
        sentence = EVIDENCE["sentences"][0]
        return {"task_id": task_id, "input_snapshot_id": source.input_snapshot_id,
                "passages": [{**sentence, "content": sentence["text"], "title": "缓存资料"}]}
    events, result = run(source, catalog=catalog, client=DeckModel(catalog.resolve(name).kind, name != "slide-deck"),
                         passage_search=search, frame_reader=lambda _: image_bytes())
    payload = result.result_payload
    assert len(events) == 5
    assert payload["pipeline_stats"]["repair_calls"] == 0
    assert not evidence_errors(payload["ir"], payload["evidence"])
    files = payload["candidate"]["required_files"]
    roles = [f["role"] for f in files]
    assert "PRIMARY_PDF" in roles if name == "lecture-notes-pdf" else "PRIMARY_PPTX" in roles
    if name != "lecture-notes-pdf":
        assert roles.count("SLIDE_PREVIEW") == 5
    else:
        assert payload["export_trace"]["status"] == "COMPILED"
    root = tmp_path / "bilibili-render-pdf/exports/pipeline-task"
    for item in files[1:]:
        assert (root / item["file_name"]).stat().st_size == item["size_bytes"]
